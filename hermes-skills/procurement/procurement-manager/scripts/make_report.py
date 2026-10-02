#!/usr/bin/env python3
"""選品調研雙報告生成器：讀調研 JSON → 同時輸出 HTML 預覽報告 + PPT 匯報文件。
用法: python3 make_report.py --json data.json --outdir <目錄> [--prefix 選品調研報告]
結構隨機應變：缺的 section 自動跳過；sections 支援 points/checklist/table/text 四型。
樣式：微軟正黑體，深青主色＋火焰橙點綴，鑄鐵黑封面（與《1688選品篩選工作冊》同族）。"""
import argparse, html, json, os, datetime
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR


def sanitize(o):
    """emoji/組合符號在服務器 PDF 渲染會變豆腐塊——遞歸替換成字體安全的符號"""
    m = {"✅": "✓", "❓": "？", "⚠️": "△", "⚠": "△", "\uFE0F": "", "☑": "✓"}
    if isinstance(o, str):
        for k, v in m.items():
            o = o.replace(k, v)
        return o
    if isinstance(o, list):
        return [sanitize(x) for x in o]
    if isinstance(o, dict):
        return {k: sanitize(v) for k, v in o.items()}
    return o

INK, PRIMARY, PRIMARY_D = RGBColor(0x23, 0x27, 0x2F), RGBColor(0x1E, 0x6E, 0x68), RGBColor(0x15, 0x51, 0x4C)
ACCENT, ACCENT_T = RGBColor(0xE8, 0x59, 0x0C), RGBColor(0xFD, 0xEE, 0xE3)
TEXT, MUTED, LINE, TINT = RGBColor(0x1F, 0x29, 0x37), RGBColor(0x64, 0x74, 0x8B), RGBColor(0xE2, 0xE8, 0xF0), RGBColor(0xEA, 0xF3, 0xF2)
F = "Microsoft JhengHei"
W = 13.333


def _tf(box, size, color, bold=False, align=PP_ALIGN.LEFT):
    box.word_wrap = True
    p = box.paragraphs[0]
    p.alignment = align
    for r in p.runs:
        r.font.name, r.font.size, r.font.color.rgb, r.font.bold = F, Pt(size), color, bold
    return p


def add_text(sl, x, y, w, h, text, size, color, bold=False, align=PP_ALIGN.LEFT, spacing=None):
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    lines = str(text).split("\n")
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if spacing:
            p.line_spacing = spacing
        r = p.add_run()
        r.text = ln
        r.font.name, r.font.size, r.font.color.rgb, r.font.bold = F, Pt(size), color, bold
    return tb


def add_rect(sl, x, y, w, h, fill, line_color=None, radius=None):
    from pptx.enum.shapes import MSO_SHAPE
    shp = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line_color:
        shp.line.color.rgb = line_color
        shp.line.width = Pt(1.25)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    if radius:
        try:
            shp.adjustments[0] = radius
        except Exception:
            pass
    return shp


def slide_title(sl, kicker, title):
    add_text(sl, 0.5, 0.32, 9, 0.3, kicker, 12, ACCENT, bold=True)
    add_text(sl, 0.5, 0.6, W - 1, 0.65, title, 26, TEXT, bold=True)


def badge(sl, x, y, verdict, scale=1.0):
    w, h = 2.0 * scale, 0.62 * scale
    color = ACCENT if verdict == "打樣" else (PRIMARY if verdict == "試單" else MUTED)
    shp = add_rect(sl, x, y, w, h, color, radius=0.5)
    tf = shp.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = verdict
    r.font.name, r.font.size, r.font.color.rgb, r.font.bold = F, Pt(18 * scale), RGBColor(0xFF, 0xFF, 0xFF), True


# ---------- 各型 section 渲染（PPT） ----------
def render_points(sl, y0, sec):
    for i, it in enumerate(sec.get("items", [])):
        y = y0 + i * 0.98
        add_text(sl, 0.6, y, 5.4, 0.4, "◆ " + it.get("head", ""), 15, PRIMARY_D, bold=True)
        add_text(sl, 0.85, y + 0.4, W - 1.5, 0.5, it.get("text", ""), 13, MUTED)
        if i < len(sec["items"]) - 1:
            ln = sl.shapes.add_shape(1, Inches(0.6), Inches(y + 0.9), Inches(W - 1.2), Emu(9525))
            ln.fill.solid()
            ln.fill.fore_color.rgb = LINE
            ln.line.fill.background()
            ln.shadow.inherit = False


def render_checklist(sl, y0, sec):
    y = y0
    for it in sec.get("items", []):
        ok = it.get("pass")
        mark, mc = ("✓", PRIMARY) if ok else ("✗", ACCENT)
        add_rect(sl, 0.6, y + 0.05, 0.42, 0.42, TINT if ok else ACCENT_T, radius=0.3)
        add_text(sl, 0.6, y + 0.07, 0.42, 0.38, mark, 16, mc, bold=True, align=PP_ALIGN.CENTER)
        add_text(sl, 1.2, y, 5.6, 0.4, it.get("text", ""), 15, TEXT, bold=True)
        add_text(sl, 1.2, y + 0.4, W - 2, 0.42, it.get("note", ""), 12.5, MUTED)
        y += 0.92


def render_table(sl, y0, sec):
    headers, rows = sec.get("headers", []), sec.get("rows", [])
    n = len(rows) + 1
    h = min(0.55 * n + 0.15, 7.05 - y0)
    tbl = sl.shapes.add_table(n, len(headers), Inches(0.5), Inches(y0), Inches(W - 1), Inches(h)).table
    total = W - 1
    widths = sec.get("colw") or [total / len(headers)] * len(headers)
    scale = total / sum(widths)
    for ci, wd in enumerate(widths):
        tbl.columns[ci].width = Emu(int(Inches(wd * scale)))
    for ci, htxt in enumerate(headers):
        c = tbl.cell(0, ci)
        c.text = str(htxt)
        _tf(c.text_frame, 13, RGBColor(0xFF, 0xFF, 0xFF), bold=True)
        c.fill.solid()
        c.fill.fore_color.rgb = PRIMARY
    for ri, row in enumerate(rows, start=1):
        for ci, val in enumerate(row):
            c = tbl.cell(ri, ci)
            c.text = str(val)
            _tf(c.text_frame, 12.5, TEXT)
            c.fill.solid()
            c.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF) if ri % 2 else RGBColor(0xF6, 0xFA, 0xF9)
    if sec.get("note"):
        add_text(sl, 0.5, min(y0 + h + 0.12, 7.0), W - 1, 0.4, sec["note"], 11.5, MUTED)


def render_text(sl, y0, sec):
    paras = sec.get("paragraphs", [])
    if len(paras) <= 5:
        # 少量段落 → 卡片行鋪滿版面，避免大面積留白
        n = len(paras)
        avail = 6.7 - y0
        cardh = min(1.3, (avail - 0.3 * (n - 1)) / n)
        for i, pp in enumerate(paras):
            y = y0 + i * (cardh + 0.3)
            add_rect(sl, 0.6, y, W - 1.2, cardh, RGBColor(0xF6, 0xFA, 0xF9), line_color=LINE, radius=0.08)
            add_text(sl, 0.95, y + cardh / 2 - 0.3, W - 1.9, cardh - 0.35, pp, 14.5, TEXT, spacing=1.3)
    else:
        add_text(sl, 0.6, y0, W - 1.2, 7.0 - y0, "\n".join(paras), 14, TEXT, spacing=1.35)


RENDER = {"points": render_points, "checklist": render_checklist, "table": render_table, "text": render_text}


def build_pptx(data, path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(7.5)
    blank = prs.slide_layouts[6]

    # 封面（鑄鐵黑）
    sl = prs.slides.add_slide(blank)
    bg = sl.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = INK
    bg.line.fill.background()
    bg.shadow.inherit = False
    add_text(sl, 0.9, 1.7, 9, 0.4, "MAINPLAN 敏寶 · 採購調研", 14, ACCENT, bold=True)
    add_text(sl, 0.9, 2.15, 11, 1.2, "選品調研報告 · " + data.get("category", ""), 44, RGBColor(0xFF, 0xFF, 0xFF), bold=True)
    add_text(sl, 0.9, 3.5, 10, 0.5, data.get("window", "") + "　|　" + data.get("date", ""), 15, RGBColor(0xB9, 0xC6, 0xD6))
    badge(sl, 0.9, 4.3, data.get("conclusion", {}).get("verdict", "待定"), 1.15)
    reasons = data.get("conclusion", {}).get("reasons", [])
    if reasons:
        add_text(sl, 0.9, 5.25, 10.5, 1.2, "結論要點：" + "；".join(reasons[:3]), 13, RGBColor(0x9F, 0xAD, 0xBE), spacing=1.3)

    # 數據說明（框高隨文字自適應，避免大塊留白）
    if data.get("data_note"):
        sl = prs.slides.add_slide(blank)
        slide_title(sl, "必讀", "數據說明與能力邊界")
        cpl = 58  # 每行約 58 個全角字符（14pt, 12.3in 寬）
        import math
        lines = max(2, math.ceil(len(data["data_note"]) / cpl))
        box_h = 0.55 + lines * 0.33
        add_rect(sl, 0.5, 1.55, W - 1, box_h, ACCENT_T, radius=0.06)
        add_text(sl, 0.85, 1.78, W - 1.7, box_h - 0.3, data["data_note"], 14, TEXT, spacing=1.35)
        y = 1.55 + box_h + 0.3
        mks = [("✓", "有來源——可點連結核實", PRIMARY), ("△", "估算值——邏輯見文字說明", ACCENT), ("？", "未能核實——不作決策依據", MUTED)]
        for mk, desc, mc in mks:
            add_rect(sl, 0.85, y + 0.02, 0.4, 0.4, TINT, radius=0.3)
            add_text(sl, 0.85, y + 0.04, 0.4, 0.36, mk, 15, mc, bold=True, align=PP_ALIGN.CENTER)
            add_text(sl, 1.45, y, 8, 0.4, desc, 14, TEXT)
            y += 0.58
        add_text(sl, 0.85, y + 0.25, W - 1.7, 1.2, "怎麼用這份報告：採購經理看「結論與行動」頁直接執行；老闆重點看趨勢與廠商頁；\n所有 △／？項在下單前人工核實。", 13, MUTED, spacing=1.4)

    # 關鍵數字
    metrics = data.get("metrics") or []
    if metrics:
        sl = prs.slides.add_slide(blank)
        slide_title(sl, "市場快照", "關鍵數字")
        n = len(metrics)
        cwid = min(2.9, (W - 1 - 0.3 * (n - 1)) / max(n, 1))
        for i, m in enumerate(metrics[:5]):
            x = 0.5 + i * (cwid + 0.3)
            add_rect(sl, x, 2.0, cwid, 2.35, RGBColor(0xFF, 0xFF, 0xFF), line_color=LINE, radius=0.08)
            v = str(m.get("value", ""))
            vsize = 30 if len(v) <= 6 else (22 if len(v) <= 10 else 17)
            add_text(sl, x + 0.25, 2.3, cwid - 0.5, 1.0, v, vsize, PRIMARY, bold=True)
            add_text(sl, x + 0.25, 3.35, cwid - 0.5, 0.45, str(m.get("label", "")), 13.5, TEXT, bold=True)
            if m.get("note"):
                add_text(sl, x + 0.25, 3.82, cwid - 0.5, 0.45, str(m["note"]), 11, MUTED)
        if data.get("data_note"):
            add_rect(sl, 0, 5.65, W, 0.8, TINT)
            add_text(sl, 0.5, 5.86, W - 1, 0.45, "口徑與可信度分級詳見「數據說明」頁；△／？項不作決策依據。", 12.5, PRIMARY_D)

    # 各 section（每個一頁；table 型項目多時拆頁由調研端控制）
    for sec in data.get("sections", []):
        sl = prs.slides.add_slide(blank)
        slide_title(sl, sec.get("kicker", ""), sec.get("title", ""))
        RENDER.get(sec.get("type"), render_text)(sl, 1.55, sec)

    # 結論與行動
    sl = prs.slides.add_slide(blank)
    slide_title(sl, "結論", "結論與下一步行動")
    badge(sl, 0.6, 1.7, data.get("conclusion", {}).get("verdict", "待定"), 1.1)
    reasons = data.get("conclusion", {}).get("reasons", [])
    for i, rz in enumerate(reasons):
        add_text(sl, 0.7, 2.65 + i * 0.55, W - 1.6, 0.45, "· " + rz, 14.5, TEXT)
    acts = data.get("actions") or []
    if acts:
        add_text(sl, 0.6, 4.6, 4, 0.4, "下一步", 15, PRIMARY_D, bold=True)
        for i, a in enumerate(acts):
            add_text(sl, 0.7, 5.05 + i * 0.5, W - 1.6, 0.45, f"{i+1}. {a}", 14, MUTED)

    prs.save(path)


# ---------- HTML ----------
def esc(s):
    return html.escape(str(s), quote=False)


def build_html(data, path):
    verdict = data.get("conclusion", {}).get("verdict", "待定")
    vcolor = "#E8590C" if verdict == "打樣" else ("#1E6E68" if verdict == "試單" else "#64748B")
    parts = ["""<!doctype html><html lang="zh-Hant"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1"><style>
*{box-sizing:border-box;margin:0}body{font-family:'Microsoft JhengHei',system-ui;background:#f4f6f9;color:#1f2937}
.wrap{max-width:980px;margin:26px auto;padding:0 14px}
.hero{background:#23272f;border-radius:14px;padding:30px 34px;color:#fff}
.hero h1{font-size:27px;margin:6px 0 4px}.hero .k{color:#e8590c;font-size:13px;font-weight:700}
.hero .m{color:#b9c6d6;font-size:14px;margin-top:6px}
.badge{display:inline-block;background:%VC%;color:#fff;border-radius:20px;padding:7px 22px;font-size:17px;font-weight:700;margin-top:14px}
.note{background:#fdeee3;border-radius:10px;padding:14px 18px;margin:14px 0;font-size:13.5px;line-height:1.7}
.cards{display:flex;gap:10px;flex-wrap:wrap;margin:14px 0}
.card{flex:1;min-width:150px;background:#fff;border-radius:10px;padding:14px 18px;box-shadow:0 1px 3px rgba(0,0,0,.08)}
.card b{display:block;font-size:25px;color:#1e6e68}.card small{color:#64748b}
.panel{background:#fff;border-radius:12px;padding:18px 22px;box-shadow:0 1px 3px rgba(0,0,0,.08);margin:14px 0}
.panel h3{font-size:16px;margin-bottom:10px;color:#1f2937}
.row{padding:9px 0;border-bottom:1px solid #eef1f5}.row:last-child{border:0}
.row .h{font-weight:700;color:#15514c;font-size:14.5px}.row .t{color:#64748b;font-size:13px;margin-top:3px;line-height:1.6}
.ok,.no{display:inline-block;width:24px;height:24px;border-radius:50%;text-align:center;line-height:24px;font-weight:700;color:#fff;margin-right:8px;font-size:14px}
.ok{background:#1e6e68}.no{background:#e8590c}
table{width:100%;border-collapse:collapse;font-size:13px}
th{background:#1e6e68;color:#fff;padding:8px 10px;text-align:left;font-weight:600}
td{padding:8px 10px;border-bottom:1px solid #eef1f5}tr:nth-child(even) td{background:#f6faf9}
.para{font-size:14px;line-height:1.8;color:#1f2937}
ul.act{list-style:none}.act li{padding:6px 0;color:#64748b;font-size:14px}
.foot{color:#8a97a8;font-size:12px;margin:16px 0;text-align:center}
</style></head><body><div class="wrap">"""]
    parts.append('<div class="hero"><div class="k">MAINPLAN 敏寶 · 採購調研</div>'
                 f'<h1>選品調研報告 · {esc(data.get("category", ""))}</h1>'
                 f'<div class="m">{esc(data.get("window",""))}　|　{esc(data.get("date",""))}</div>'
                 f'<div class="badge">結論：{esc(verdict)}</div></div>')
    if data.get("data_note"):
        parts.append(f'<div class="note"><b>數據說明：</b>{esc(data["data_note"])}</div>')
    if data.get("metrics"):
        parts.append('<div class="cards">' + "".join(
            f'<div class="card"><b>{esc(m.get("value",""))}</b><small>{esc(m.get("label",""))}'
            + (f'　{esc(m["note"])}' if m.get("note") else '') + '</small></div>' for m in data["metrics"]) + '</div>')
    for sec in data.get("sections", []):
        parts.append(f'<div class="panel"><h3>{esc(sec.get("title",""))}</h3>')
        if sec["type"] == "points":
            parts.append("".join(f'<div class="row"><div class="h">◆ {esc(it.get("head",""))}</div>'
                                 f'<div class="t">{esc(it.get("text",""))}</div></div>' for it in sec.get("items", [])))
        elif sec["type"] == "checklist":
            parts.append("".join(f'<div class="row"><span class="{"ok" if it.get("pass") else "no"}">'
                                 f'{"✓" if it.get("pass") else "✗"}</span><b>{esc(it.get("text",""))}</b>'
                                 f'<div class="t">{esc(it.get("note",""))}</div></div>' for it in sec.get("items", [])))
        elif sec["type"] == "table":
            parts.append('<table><tr>' + "".join(f'<th>{esc(h)}</th>' for h in sec.get("headers", [])) + '</tr>'
                         + "".join('<tr>' + "".join(f'<td>{esc(c)}</td>' for c in r) + '</tr>' for r in sec.get("rows", []))
                         + '</table>')
            if sec.get("note"):
                parts.append(f'<div class="t" style="margin-top:8px">{esc(sec["note"])}</div>')
        else:
            parts.append("".join(f'<p class="para">{esc(pp)}</p>' for pp in sec.get("paragraphs", [])))
        parts.append('</div>')
    reasons = data.get("conclusion", {}).get("reasons", [])
    acts = data.get("actions") or []
    parts.append('<div class="panel"><h3>結論與下一步行動</h3>'
                 + "".join(f'<p class="para">· {esc(rz)}</p>' for rz in reasons)
                 + '<ul class="act">' + "".join(f'<li>{i+1}. {esc(a)}</li>' for i, a in enumerate(acts)) + '</ul></div>')
    parts.append(f'<div class="foot">Luma · 敏寶團隊 AI 生成　|　{esc(data.get("date",""))}　|　本報告由 AI 整理，重大決策請以人工核實為準</div></div></body></html>')
    out = "\n".join(parts).replace("%VC%", vcolor)
    open(path, "w", encoding="utf-8").write(out)


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--json", required=True)
    ap.add_argument("--outdir", required=True)
    ap.add_argument("--prefix", default="選品調研報告")
    a = ap.parse_args()
    data = sanitize(json.load(open(a.json, encoding="utf-8")))
    os.makedirs(a.outdir, exist_ok=True)
    stem = f'{a.prefix}-{data.get("category","報告")}'
    html_path = os.path.join(a.outdir, stem + ".html")
    pptx_path = os.path.join(a.outdir, stem + ".pptx")
    build_html(data, html_path)
    build_pptx(data, pptx_path)
    print("HTML:", html_path)
    print("PPTX:", pptx_path)
